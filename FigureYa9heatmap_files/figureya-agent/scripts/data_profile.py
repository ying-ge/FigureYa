"""Bounded, local-only tabular profiling; semantic roles remain hypotheses."""
import collections
import csv
import datetime
import itertools
import math
import re
from pathlib import Path

MISSING = {'', 'na', 'n/a', 'nan', 'null', 'none'}
ROLE_PATTERNS = {
    'id': r'^(id|sample[_ .-]?id|subject[_ .-]?id|patient[_ .-]?id|样本|样本编号|受试者|患者编号)$',
    'group': r'^(group|condition|treatment|class|cohort|分组|组别)$',
    'time': r'^(time|os[._ ]?time|pfs[._ ]?time|followup|survival[._ ]?time|随访时间|生存时间)$',
    'event': r'^(event|status|os[._ ]?status|pfs[._ ]?status|censor|事件|结局状态)$',
    'logfc': r'^(log2?fc|log2foldchange|log[._]?fold[._]?change)$',
    'pvalue': r'^(p[._ ]?value|pval|p)$',
    'padj': r'^(padj|adj[._ ]?p[._ ]?val(?:ue)?|fdr|q[._ ]?value)$',
    'feature': r'^(gene|gene[._ ]?id|gene[._ ]?name|symbol|gsym|feature|基因)$',
    'term': r'^(term|pathway|description|通路)$',
    'effect': r'^(estimate|effect|hr|or|rr|coef)$',
    'lower': r'^(lower|lcl|ci[._]?low|lower95|conf[._]?low)$',
    'upper': r'^(upper|ucl|ci[._]?high|upper95|conf[._]?high)$',
    'label': r'^(label|outcome|truth|response|diagnosis|y|结局)$',
    'score': r'^(score|risk|probability|prob|prediction|pred|risk[._]?score)$',
    'source': r'^(source|from|源)$',
    'target': r'^(target|to|目标)$',
    'weight': r'^(weight|count|frequency|freq|n|权重|计数)$',
    'chromosome': r'^(chr|chromosome|染色体)$',
    'position': r'^(pos|position|start|位置)$',
    'donor': r'^(donor|donor_id|patient|subject)$',
    'cell': r'^(cell|cell_id|barcode)$',
}


def text_value(x):
    if x is None:
        return ''
    if isinstance(x, (datetime.datetime, datetime.date)):
        return x.isoformat()
    return str(x).strip()


def number(x):
    try:
        return float(x)
    except (ValueError, TypeError):
        return None


def quantile(values, p):
    i = (len(values) - 1) * p
    a = int(i)
    b = min(a + 1, len(values) - 1)
    return values[a] * (b - i) + values[b] * (i - a) if b != a else values[a]


def summarize(headers, rows, sampled, source, extra=None):
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Column headers must be non-empty and unique; fix/export headers explicitly')
    if not rows:
        raise ValueError('No data rows')
    if any(len(r) != len(headers) for r in rows):
        raise ValueError('Ragged table: row widths differ from header; specify the correct delimiter')
    columns = []
    roles = collections.defaultdict(list)
    warnings = []
    for idx, name in enumerate(headers):
        vals = [r[idx] for r in rows]
        present = [v for v in vals if v.lower() not in MISSING]
        numeric = [number(v) for v in present]
        finite = sorted(v for v in numeric if v is not None and math.isfinite(v))
        nonfinite = sum(v is not None and not math.isfinite(v) for v in numeric)
        unique = len(set(present))
        is_numeric = bool(present) and all(v is not None for v in numeric)
        typ = 'numeric' if is_numeric else 'text'
        if not present:
            typ = 'empty'
        col = {'name': name, 'kind': typ, 'missing': len(vals) - len(present),
               'missing_fraction': round((len(vals) - len(present)) / len(vals), 5),
               'distinct_observed': unique, 'nonfinite': nonfinite,
               'possible_identifier': bool(present) and unique == len(present),
               'possible_category': 1 < unique <= min(30, max(2, len(rows) // 5))}
        if is_numeric and finite:
            q1, median, q3 = [quantile(finite, p) for p in (.25, .5, .75)]
            iqr = q3 - q1
            col.update(min=finite[0], max=finite[-1], median=median, q1=q1, q3=q3,
                       zeros=sum(v == 0 for v in finite), negative=sum(v < 0 for v in finite),
                       integer_like=all(v.is_integer() for v in finite),
                       binary_01=set(finite) <= {0, 1}, constant=len(set(finite)) == 1,
                       tukey_outliers=sum(v < q1 - 1.5 * iqr or v > q3 + 1.5 * iqr for v in finite))
        # Counts only: do not echo identifiers or patient values into the report.
        if col['possible_category']:
            col['category_counts_sorted'] = sorted(collections.Counter(present).values(), reverse=True)
        if nonfinite:
            warnings.append(f'{name}: contains {nonfinite} infinite values')
        for role, pattern in ROLE_PATTERNS.items():
            if re.search(pattern, name, re.I):
                roles[role].append(name)
        columns.append(col)
    numeric_columns = [c['name'] for c in columns if c['kind'] == 'numeric']
    category_columns = [c['name'] for c in columns if c['possible_category']]
    if sampled:
        warnings.append('Prefix sample only; tail, rare categories and total row count are unknown. Validate full selected columns before plotting.')
    warnings.append('Names and numeric codes do not establish units, count normalization, event meaning, pairing, or biological independence.')
    return {'schema_version': 1, 'source': str(source), 'rows_observed': len(rows), 'sampled': sampled,
            'row_count_exact': None if sampled else len(rows), 'sampling': 'first rows',
            'columns': columns, 'numeric_columns': numeric_columns, 'category_candidates': category_columns,
            'role_candidates': dict(roles), 'matrix_candidate': len(numeric_columns) >= 3,
            'duplicate_rows_observed': len(rows) - len(set(tuple(r) for r in rows)),
            'missing_tokens': sorted(MISSING), 'warnings': warnings, **(extra or {})}


def profile(path, limit=10000, delimiter=None, encoding='utf-8-sig', sheet=None):
    path = Path(path).resolve()
    if limit < 2 or limit > 100000:
        raise ValueError('Sample limit must be between 2 and 100000')
    if path.suffix.lower() in {'.xlsx', '.xlsm'}:
        try:
            import openpyxl
        except ImportError as e:
            raise ValueError('XLSX needs openpyxl; use the bundled runtime or export CSV') from e
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError('Multiple worksheets; select --sheet explicitly: ' + ', '.join(book.sheetnames))
            name = sheet or book.sheetnames[0]
            if name not in book.sheetnames:
                raise ValueError('Worksheet not found: ' + name)
            iterator = book[name].iter_rows(values_only=True)
            headers = [text_value(x) for x in next(iterator)]
            rows = [[text_value(v) for v in r] for r in itertools.islice(iterator, limit + 1)]
            result = summarize(headers, rows[:limit], len(rows) > limit, path, {'sheet': name})
            result['warnings'].append('XLSX reads cached formula values; absent/stale caches require recalculation in the spreadsheet application.')
            return result
        finally:
            book.close()
    if path.suffix.lower() not in {'.csv', '.tsv', '.txt'}:
        raise ValueError('Supported: CSV/TSV/TXT/XLSX. RDS/H5AD/MTX need a format-aware reader; do not flatten blindly.')
    with path.open(encoding=encoding, newline='') as handle:
        if delimiter == 'tab':
            delimiter = '\t'
        if delimiter == 'whitespace':
            iterator = (re.split(r'\s+', line.strip()) for line in handle if line.strip())
        else:
            if delimiter is None:
                if path.suffix.lower() == '.tsv':
                    delimiter = '\t'
                else:
                    sample = handle.read(65536)
                    handle.seek(0)
                    try:
                        delimiter = csv.Sniffer().sniff(sample, delimiters=',\t;|').delimiter
                    except csv.Error as e:
                        raise ValueError('Cannot infer delimiter. Use --delimiter tab/comma character/whitespace explicitly.') from e
            iterator = csv.reader(handle, delimiter=delimiter)
        try:
            headers = [text_value(x) for x in next(iterator)]
        except StopIteration as e:
            raise ValueError('Empty file') from e
        rows = [[text_value(v) for v in r] for r in itertools.islice(iterator, limit + 1)]
    return summarize(headers, rows[:limit], len(rows) > limit, path, {'delimiter': delimiter, 'encoding': encoding})
